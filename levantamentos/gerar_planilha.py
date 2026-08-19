# -*- coding: utf-8 -*-
"""Gera o workbook de Levantamento Ambiental das Fazendas Campanelli."""
import json, os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

BASE = os.path.dirname(os.path.abspath(__file__))
DATA = json.load(open(os.path.join(BASE, 'dados_base.json')))
OUT = os.environ.get('OUT', os.path.join(BASE, 'out.xlsx'))

AZUL   = '024168'   # azul institucional Vergo
VERDE  = '0D7B3C'   # verde institucional Vergo
CINZA  = 'F2F5F7'
AMARELO= 'FFF4CE'

thin = Side(style='thin', color='BFC9D1')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = openpyxl.Workbook()
wb.remove(wb.active)

STATUS_LEV = '"Não iniciado,Em consulta,Concluído,Pendência documental,Não se aplica"'


def header(ws, cols, row=1, fill=AZUL):
    for i, (title, width) in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=title)
        c.font = Font(bold=True, color='FFFFFF', size=10)
        c.fill = PatternFill('solid', fgColor=fill)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[row].height = 42
    ws.freeze_panes = ws.cell(row=row + 1, column=1)
    ws.auto_filter.ref = f"A{row}:{get_column_letter(len(cols))}{row}"


def body(ws, first_row, last_row, ncols):
    for r in range(first_row, last_row + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            cell.font = Font(size=10)
        if r % 2 == 0:
            for c in range(1, ncols + 1):
                ws.cell(row=r, column=c).fill = PatternFill('solid', fgColor=CINZA)


def dv(ws, itens, col, first, last):
    """Cria validacao de lista apoiada na aba oculta 'Listas'.

    O Excel limita formula1 embutida a 255 caracteres; listas longas precisam
    viver numa faixa de celulas."""
    if isinstance(itens, str):
        itens = [i for i in itens.strip('"').split(',')]
    key = '|'.join(itens)
    ref = LIST_REFS.get(key)
    if ref is None:
        c = LIST_STATE['col'] = LIST_STATE['col'] + 1
        letter = get_column_letter(c)
        LISTAS.cell(row=1, column=c, value=f'lista_{c}').font = Font(bold=True, size=9)
        for i, v in enumerate(itens, start=2):
            LISTAS.cell(row=i, column=c, value=v)
        ref = f"=Listas!${letter}$2:${letter}${len(itens) + 1}"
        LIST_REFS[key] = ref
    v = DataValidation(type='list', formula1=ref, allow_blank=True, showDropDown=False)
    ws.add_data_validation(v)
    v.add(f'{col}{first}:{col}{last}')


LISTAS = wb.create_sheet('Listas')
LISTAS.sheet_state = 'hidden'
LIST_REFS = {}
LIST_STATE = {'col': 0}

PROP = DATA['proprias']
PARC = DATA['parceiras']
IMOVEIS = [(p['imovel'], p['municipio'], 'Própria') for p in PROP] + \
          [(p['imovel'], '', 'Parceira') for p in PARC]
LISTA_IMOVEIS = [i[0] for i in IMOVEIS]

# ---------------------------------------------------------------- 1. Instruções
ws = wb.create_sheet('Instruções')
ws.sheet_view.showGridLines = False
ws['A1'] = 'LEVANTAMENTO AMBIENTAL — FAZENDAS CAMPANELLI'
ws['A1'].font = Font(bold=True, size=16, color=AZUL)
ws['A2'] = 'Roteiro de consulta e instrumento de coleta — CAR, Outorgas, Autos de Infração/TCRA, Processos CETESB, Incêndios e Licenciamentos'
ws['A2'].font = Font(size=10, color='444444')
ws['A4'] = ('ATENÇÃO: as abas de dados estão ESTRUTURADAS E PRÉ-PREENCHIDAS com o que consta na planilha de origem '
            '(imóvel, município, CCIR, matrícula e área). Os campos ambientais estão EM BRANCO e devem ser preenchidos '
            'a partir das consultas oficiais indicadas abaixo — nenhum dado ambiental foi presumido ou estimado.')
ws['A4'].font = Font(size=10, bold=True, color='9C5700')
ws['A4'].fill = PatternFill('solid', fgColor=AMARELO)
ws['A4'].alignment = Alignment(wrap_text=True, vertical='center')
ws.merge_cells('A4:F4')
ws.row_dimensions[4].height = 46

cols = [('Item', 26), ('Órgão / Sistema', 26), ('Onde consultar (portal)', 42),
        ('Dado necessário para a consulta', 34), ('Documento a obter', 32), ('Aba de destino', 18)]
header(ws, cols, row=6)
fontes = [
 ('N° do CAR', 'SICAR / SFB — Consulta Pública',
  'consultapublica.car.gov.br/publico/imoveis/index',
  'N° do recibo do CAR ou CPF/CNPJ do declarante; alternativamente o polígono do imóvel',
  'Demonstrativo do CAR + Recibo de inscrição', 'CAR'),
 ('N° do CAR (SP)', 'SICAR-SP / SIGAM — SEMIL-SP',
  'sigam.ambiente.sp.gov.br  •  car.sp.gov.br',
  'CPF/CNPJ do proprietário ou n° do CAR',
  'Situação da inscrição, análise e adesão ao PRA', 'CAR'),
 ('Outorga (águas estaduais)', 'DAEE-SP',
  'outorga.daee.sp.gov.br  •  aplicacoes.daee.sp.gov.br',
  'CPF/CNPJ do requerente ou n° do processo/portaria DAEE',
  'Portaria de outorga / Licença de execução / Dispensa', 'Outorgas'),
 ('Outorga (águas federais)', 'ANA — REGLA / SNIRH',
  'gov.br/ana  •  snirh.gov.br',
  'CPF/CNPJ ou n° da resolução. Aplica-se a captações no Rio Grande, Paranapanema e Paraná (rios de domínio da União)',
  'Resolução de outorga ANA', 'Outorgas'),
 ('Poços tubulares', 'DAEE / SIAGAS-CPRM',
  'siagasweb.sgb.gov.br',
  'Município + coordenadas do poço',
  'Cadastro do poço e outorga de perfuração/uso', 'Outorgas'),
 ('Autos de Infração (estadual)', 'CETESB / Polícia Militar Ambiental',
  'Agência Ambiental da CETESB do município  •  cetesb.sp.gov.br',
  'CPF/CNPJ do autuado e/ou n° do auto (AIA/AIPI)',
  'Auto de Infração, Auto de Imposição de Penalidade, extrato de débito', 'Autos e TCRA'),
 ('Autos de Infração (federal)', 'IBAMA — Consulta de autuações',
  'servicos.ibama.gov.br  •  Painel de Autos de Infração (dadosabertos.ibama.gov.br)',
  'CPF/CNPJ do autuado',
  'Auto de infração e situação do débito/embargo', 'Autos e TCRA'),
 ('TCRA / TAC', 'CETESB / SEMIL-SP / Ministério Público',
  'Agência Ambiental da CETESB  •  MPSP (comarca do imóvel)',
  'N° do processo administrativo ou CPF/CNPJ',
  'Termo de Compromisso de Recuperação Ambiental e PRAD vinculado', 'Autos e TCRA'),
 ('Processo CETESB', 'CETESB — Agência Ambiental',
  'e-ambiente / SIGAM  •  protocolo na Agência Ambiental',
  'N° do processo ou CPF/CNPJ do interessado',
  'Extrato/andamento do processo administrativo', 'Processos CETESB'),
 ('Incêndios / queimadas', 'INPE — Programa Queimadas (BDQueimadas)',
  'terrabrasilis.dpi.inpe.br/queimadas  •  queimadas.dgi.inpe.br',
  'Polígono do imóvel (shapefile do CAR) ou município + período',
  'Relatório de focos de calor e área queimada por data', 'Incêndios'),
 ('Incêndios — ocorrência', 'Corpo de Bombeiros / Defesa Civil / PM Ambiental',
  'Unidade da comarca do imóvel',
  'Data e local da ocorrência',
  'Boletim de Ocorrência e eventual auto de infração vinculado', 'Incêndios'),
 ('Licenciamento ambiental', 'CETESB — Licenciamento',
  'licenciamento.cetesb.sp.gov.br',
  'CPF/CNPJ do empreendedor ou n° da licença/processo',
  'LP / LI / LO / Licença Simplificada / CADRI e condicionantes', 'Licenciamentos'),
 ('Supressão de vegetação e APP', 'SEMIL-SP / CFB — via SIGAM',
  'sigam.ambiente.sp.gov.br',
  'N° do processo ou CPF/CNPJ',
  'Autorização de supressão / intervenção em APP e compensação', 'Licenciamentos'),
 ('Licença de uso de irrigação', 'CATI/SAA-SP e CETESB',
  'Conforme o porte e a atividade do empreendimento',
  'Descrição da atividade e área irrigada',
  'Licença ou declaração de dispensa', 'Licenciamentos'),
]
r = 7
for f in fontes:
    for i, v in enumerate(f, start=1):
        ws.cell(row=r, column=i, value=v)
    r += 1
body(ws, 7, r - 1, len(cols))

ws.cell(row=r + 1, column=1, value='CAMINHO MAIS CURTO: a maior parte das consultas acima é feita por CPF/CNPJ do proprietário. '
        'Levantando por CNPJ (SICAR, DAEE, IBAMA, CETESB) obtém-se de uma só vez a lista completa de registros do grupo, '
        'que depois é rateada por imóvel nas abas deste arquivo.').font = Font(size=10, italic=True, color=VERDE)
ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=6)
ws.cell(row=r + 1, column=1).alignment = Alignment(wrap_text=True, vertical='center')
ws.row_dimensions[r + 1].height = 32

r += 3
ws.cell(row=r, column=1, value='MUNICÍPIOS DO PORTFÓLIO — REFERÊNCIA PARA PROTOCOLO (confirmar em cetesb.sp.gov.br/agencias-ambientais)').font = Font(bold=True, size=11, color=AZUL)
r += 1
header(ws, [('Município', 26), ('Imóveis', 46), ('Agência Ambiental CETESB (a confirmar)', 26), ('UGRHI / Bacia (a confirmar)', 34), ('', 2), ('', 2)], row=r)
ws.freeze_panes = None
ws.auto_filter.ref = None
muns = {}
for p in PROP:
    muns.setdefault(p['municipio'], []).append(p['imovel'])
r += 1
first_m = r
for m, lst in sorted(muns.items()):
    ws.cell(row=r, column=1, value=m)
    ws.cell(row=r, column=2, value=', '.join(lst))
    r += 1
body(ws, first_m, r - 1, 4)

# ------------------------------------------------------- 2. Resumo por Imóvel
ws = wb.create_sheet('Resumo por Imóvel')
cols = [('Imóvel', 34), ('Município', 22), ('Tipo', 11), ('CCIR', 20), ('Matrículas', 26),
        ('Área Mat. (ha)', 13), ('Área Imóvel (ha)', 14), ('N° do CAR (principal)', 26),
        ('Situação do CAR', 15), ('Outorgas (qtd)', 11), ('Autos de Infração (qtd)', 12),
        ('TCRA / TAC (qtd)', 10), ('Processos CETESB (qtd)', 12), ('Incêndios (qtd)', 11),
        ('Licenciamentos (qtd)', 12), ('Status do levantamento', 20), ('Responsável', 18),
        ('Data da consulta', 14), ('Observações', 44)]
header(ws, cols)
r = 2
for p in PROP:
    mats = ', '.join(str(m['matricula']) for m in p['matriculas'])
    area = sum(m['area'] or 0 for m in p['matriculas'])
    ws.cell(row=r, column=1, value=p['imovel'])
    ws.cell(row=r, column=2, value=p['municipio'])
    ws.cell(row=r, column=3, value='Própria')
    ws.cell(row=r, column=4, value=p['ccir'])
    ws.cell(row=r, column=5, value=mats)
    ws.cell(row=r, column=6, value=round(area, 4)).number_format = '#,##0.0000'
    ws.cell(row=r, column=7, value=round(p['area_imovel'], 4) if p['area_imovel'] else None).number_format = '#,##0.0000'
    r += 1
for p in PARC:
    ws.cell(row=r, column=1, value=p['imovel'])
    ws.cell(row=r, column=3, value='Parceira')
    ws.cell(row=r, column=4, value=p.get('incra'))
    ws.cell(row=r, column=5, value=', '.join(p['matriculas']))
    r += 1
last = r - 1
for rr in range(2, last + 1):
    ws.cell(row=rr, column=10,  value=f"=COUNTIF(Outorgas!$A:$A,$A{rr})")
    ws.cell(row=rr, column=11, value=(
        f'=COUNTIFS(\'Autos e TCRA\'!$A:$A,$A{rr},\'Autos e TCRA\'!$D:$D,"Auto de Infração*")'
        f'+COUNTIFS(\'Autos e TCRA\'!$A:$A,$A{rr},\'Autos e TCRA\'!$D:$D,"Auto de Imposição*")'))
    ws.cell(row=rr, column=12, value=(
        f'=COUNTIFS(\'Autos e TCRA\'!$A:$A,$A{rr},\'Autos e TCRA\'!$D:$D,"TCRA")'
        f'+COUNTIFS(\'Autos e TCRA\'!$A:$A,$A{rr},\'Autos e TCRA\'!$D:$D,"TAC")'))
    ws.cell(row=rr, column=13, value=f"=COUNTIF('Processos CETESB'!$A:$A,$A{rr})")
    ws.cell(row=rr, column=14, value=f"=COUNTIF(Incêndios!$A:$A,$A{rr})")
    ws.cell(row=rr, column=15, value=f"=COUNTIF(Licenciamentos!$A:$A,$A{rr})")
    ws.cell(row=rr, column=16, value='Não iniciado')
    for cc in range(10, 16):
        ws.cell(row=rr, column=cc).alignment = Alignment(horizontal='center', vertical='center')
body(ws, 2, last, len(cols))
for rr in range(2, last + 1):
    for cc in range(8, 10):
        ws.cell(row=rr, column=cc).fill = PatternFill('solid', fgColor=AMARELO)
dv(ws, '"Ativo,Pendente,Em análise,Cancelado,Suspenso,Não inscrito,A verificar"', 'I', 2, last)
dv(ws, STATUS_LEV, 'P', 2, last)
ws.cell(row=last + 2, column=1, value='Campos em amarelo = preenchimento manual a partir da consulta oficial. Colunas J:O são contagens automáticas das abas de detalhe.').font = Font(size=9, italic=True, color='666666')

# ------------------------------------------------------------------- 3. CAR
ws = wb.create_sheet('CAR')
cols = [('Imóvel', 32), ('Município', 20), ('Matrícula', 12), ('Área da matrícula (ha)', 14),
        ('N° do CAR (recibo)', 34), ('Situação do CAR', 16), ('Data de inscrição', 14),
        ('Última retificação', 14), ('Área do imóvel no CAR (ha)', 15), ('Reserva Legal declarada (ha)', 15),
        ('RL averbada / compensada', 18), ('APP (ha)', 11), ('Área consolidada (ha)', 14),
        ('Passivo / área a recuperar (ha)', 16), ('Adesão ao PRA', 14), ('Sobreposições identificadas', 26),
        ('Fonte / data da consulta', 20), ('Observações', 40)]
header(ws, cols)
r = 2
for p in PROP:
    for m in p['matriculas']:
        ws.cell(row=r, column=1, value=p['imovel'])
        ws.cell(row=r, column=2, value=p['municipio'])
        ws.cell(row=r, column=3, value=m['matricula'])
        ws.cell(row=r, column=4, value=m['area']).number_format = '#,##0.0000'
        r += 1
for p in PARC:
    for m in p['matriculas']:
        ws.cell(row=r, column=1, value=p['imovel'])
        ws.cell(row=r, column=3, value=m)
        r += 1
last = r - 1
body(ws, 2, last, len(cols))
for rr in range(2, last + 1):
    for cc in range(5, 17):
        ws.cell(row=rr, column=cc).fill = PatternFill('solid', fgColor=AMARELO)
dv(ws, '"Ativo,Pendente,Em análise,Analisado sem pendências,Cancelado,Suspenso,Não inscrito"', 'F', 2, last)
dv(ws, '"Sim,Não,Em processo,Não se aplica"', 'K', 2, last)
dv(ws, '"Sim,Não,Não se aplica,A verificar"', 'O', 2, last)

# -------------------------------------------------------------- 4. Outorgas
ws = wb.create_sheet('Outorgas')
cols = [('Imóvel', 32), ('Município', 20), ('Matrícula', 12), ('Tipo de uso', 24),
        ('Corpo hídrico / aquífero', 24), ('Órgão outorgante', 16), ('N° da portaria / resolução', 22),
        ('N° do processo', 18), ('Vazão outorgada (m³/h)', 14), ('Regime (h/dia — dias/mês)', 16),
        ('Volume anual (m³)', 14), ('Data de emissão', 14), ('Validade', 14), ('Situação', 16),
        ('Coordenadas do ponto', 22), ('UGRHI / bacia', 18), ('Fonte / data da consulta', 20), ('Observações', 40)]
header(ws, cols)
last = 61
body(ws, 2, last, len(cols))
dv(ws, LISTA_IMOVEIS, 'A', 2, last)
dv(ws, '"Captação superficial,Poço tubular profundo,Poço raso/cacimba,Barramento/reservatório,Travessia,Lançamento de efluente,Dispensa de outorga,Uso insignificante"', 'D', 2, last)
dv(ws, '"DAEE-SP,ANA,Município,Outro"', 'F', 2, last)
dv(ws, '"Vigente,Vencida,Em renovação,Em análise,Cancelada,Não possui,A verificar"', 'N', 2, last)
for rr in range(2, last + 1):
    ws.cell(row=rr, column=12).number_format = 'DD/MM/YYYY'
    ws.cell(row=rr, column=13).number_format = 'DD/MM/YYYY'

# ---------------------------------------------------------- 5. Autos e TCRA
ws = wb.create_sheet('Autos e TCRA')
cols = [('Imóvel', 32), ('Município', 20), ('Matrícula', 12), ('Tipo de documento', 26),
        ('N° do documento', 20), ('Órgão', 20), ('Data', 13), ('Descrição da infração / objeto', 46),
        ('Enquadramento legal', 24), ('Valor da multa (R$)', 15), ('Situação', 20),
        ('TCRA / TAC vinculado', 20), ('Área a recuperar (ha)', 14), ('PRAD protocolado', 15),
        ('Prazo de execução', 14), ('Status de cumprimento', 20), ('Fonte / data da consulta', 20), ('Observações', 40)]
header(ws, cols)
last = 61
body(ws, 2, last, len(cols))
dv(ws, LISTA_IMOVEIS, 'A', 2, last)
dv(ws, '"Auto de Infração Ambiental (AIA),Auto de Imposição de Penalidade,Auto de Inspeção,TCRA,TAC,Embargo,Notificação"', 'D', 2, last)
dv(ws, '"CETESB,Polícia Militar Ambiental,IBAMA,SEMIL-SP,Ministério Público,Prefeitura,Outro"', 'F', 2, last)
dv(ws, '"Em defesa,Em recurso,Aguardando julgamento,Pago,Parcelado,Convertido em serviços,Anulado,Inscrito em dívida ativa,Encerrado"', 'K', 2, last)
dv(ws, '"Sim,Não,Em elaboração,Não se aplica"', 'N', 2, last)
dv(ws, '"Não iniciado,Em execução,Concluído,Em atraso,Não se aplica"', 'P', 2, last)
for rr in range(2, last + 1):
    ws.cell(row=rr, column=7).number_format = 'DD/MM/YYYY'
    ws.cell(row=rr, column=10).number_format = 'R$ #,##0.00'
    ws.cell(row=rr, column=15).number_format = 'DD/MM/YYYY'

# ------------------------------------------------------ 6. Processos CETESB
ws = wb.create_sheet('Processos CETESB')
cols = [('Imóvel', 32), ('Município', 20), ('N° do processo CETESB', 24), ('Agência Ambiental', 24),
        ('Assunto / objeto', 46), ('Data de abertura', 15), ('Situação', 20),
        ('Última movimentação', 15), ('Exigências pendentes', 40), ('Prazo de atendimento', 15),
        ('Responsável técnico / ART', 24), ('Fonte / data da consulta', 20), ('Observações', 40)]
header(ws, cols)
last = 61
body(ws, 2, last, len(cols))
dv(ws, LISTA_IMOVEIS, 'A', 2, last)
dv(ws, '"Em análise,Em exigência,Deferido,Indeferido,Arquivado,Suspenso,Recurso"', 'G', 2, last)
for rr in range(2, last + 1):
    for cc in (6, 8, 10):
        ws.cell(row=rr, column=cc).number_format = 'DD/MM/YYYY'

# ------------------------------------------------------------- 7. Incêndios
ws = wb.create_sheet('Incêndios')
cols = [('Imóvel', 32), ('Município', 20), ('Matrícula', 12), ('Data da ocorrência', 15),
        ('Área atingida (ha)', 14), ('Cobertura atingida', 26), ('Atingiu APP/Reserva Legal', 18),
        ('Origem provável', 22), ('Focos INPE (n°/data)', 18), ('Boletim de Ocorrência', 18),
        ('Houve autuação (n° do auto)', 22), ('Ação de recuperação adotada', 34),
        ('Custo / prejuízo estimado (R$)', 16), ('Fonte / data da consulta', 20), ('Observações', 40)]
header(ws, cols)
last = 61
body(ws, 2, last, len(cols))
dv(ws, LISTA_IMOVEIS, 'A', 2, last)
dv(ws, '"Cana-de-açúcar,Pastagem,Reflorestamento,Fragmento nativo,APP,Reserva Legal,Aceiro/carreador,Misto"', 'F', 2, last)
dv(ws, '"Sim — APP,Sim — Reserva Legal,Sim — ambos,Não,A verificar"', 'G', 2, last)
dv(ws, '"Acidental,Criminosa/dolosa,Origem externa (vizinho/rodovia/ferrovia),Rede elétrica,Balão,Queima controlada autorizada,Desconhecida"', 'H', 2, last)
dv(ws, '"Sim,Não,A verificar"', 'J', 2, last)
for rr in range(2, last + 1):
    ws.cell(row=rr, column=4).number_format = 'DD/MM/YYYY'
    ws.cell(row=rr, column=13).number_format = 'R$ #,##0.00'

# -------------------------------------------------------- 8. Licenciamentos
ws = wb.create_sheet('Licenciamentos')
cols = [('Imóvel', 32), ('Município', 20), ('Matrícula', 12), ('Tipo de documento', 30),
        ('N° da licença / documento', 22), ('N° do processo', 18), ('Órgão emissor', 20),
        ('Objeto / atividade licenciada', 40), ('Data de emissão', 14), ('Validade', 14),
        ('Situação', 18), ('Prazo p/ renovação (120 dias antes)', 18), ('Condicionantes pendentes', 40),
        ('Responsável técnico / ART', 22), ('Fonte / data da consulta', 20), ('Observações', 40)]
header(ws, cols)
last = 61
body(ws, 2, last, len(cols))
dv(ws, LISTA_IMOVEIS, 'A', 2, last)
dv(ws, '"Licença Prévia (LP),Licença de Instalação (LI),Licença de Operação (LO),Licença Ambiental Simplificada,Renovação de LO,CADRI,Certificado de Dispensa de Licença,Autorização de supressão de vegetação,Autorização de intervenção em APP,Licença de queima controlada,Cadastro de fonte de poluição"', 'D', 2, last)
dv(ws, '"CETESB,SEMIL-SP,IBAMA,Prefeitura,DAEE,Outro"', 'G', 2, last)
dv(ws, '"Vigente,Vencida,Em renovação,Em análise,Cancelada,Não se aplica,A verificar"', 'K', 2, last)
for rr in range(2, last + 1):
    ws.cell(row=rr, column=9).number_format = 'DD/MM/YYYY'
    ws.cell(row=rr, column=10).number_format = 'DD/MM/YYYY'
    ws.cell(row=rr, column=12, value=f'=IF(J{rr}="","",J{rr}-120)').number_format = 'DD/MM/YYYY'

# ------------------------------------------------------------ 9/10. Base original
ORIG = os.environ.get('ORIGEM', '/root/.claude/uploads/d6e29332-3163-5b75-85ff-70728a93c9b8/'
                      '90937ec2-C_pia_de_Planilha_de_Fazendas_Campanelli1.xlsx')
src = openpyxl.load_workbook(ORIG, data_only=True) if os.path.exists(ORIG) else None
for name, dest in ([] if src is None else [('Fazendas Proprias', 'Base - Próprias'), ('Fazendas Parceiras', 'Base - Parceiras')]):
    s = src[name]
    d = wb.create_sheet(dest)
    for row in s.iter_rows():
        for c in row:
            if c.value is not None:
                d.cell(row=c.row, column=c.column, value=c.value)
    for i in range(1, s.max_column + 1):
        d.column_dimensions[get_column_letter(i)].width = 26
    d.sheet_properties.tabColor = '999999'

for name, color in (('Instruções', VERDE), ('Resumo por Imóvel', AZUL)):
    wb[name].sheet_properties.tabColor = color
for name in ('CAR', 'Outorgas', 'Autos e TCRA', 'Processos CETESB', 'Incêndios', 'Licenciamentos'):
    wb[name].sheet_properties.tabColor = '4A90A4'

wb.move_sheet('Listas', offset=len(wb.sheetnames))
wb.active = wb.sheetnames.index('Resumo por Imóvel')
wb.save(OUT)
print('gravado:', OUT)
